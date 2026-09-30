import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from functools import lru_cache
from io import BytesIO
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt


BASE = "https://learn.microsoft.com"
REPO_ROOT = Path(__file__).resolve().parent.parent
LOGO_DIR = REPO_ROOT / "assets" / "logos"
PRODUCT_ICON_DIR = REPO_ROOT / "assets" / "product-icons"
DOCUMENTS_DIR = REPO_ROOT / "assets" / "documents"
DATA_DIR = REPO_ROOT / "data"
DATA_PATH = DATA_DIR / "connectors.json"
CHANGE_LOG_PATH = DATA_DIR / "change-log.json"
PPTX_PATH = DOCUMENTS_DIR / "Power_Platform_Gov_Connectors.pptx"
DOCX_PATH = DOCUMENTS_DIR / "Power_Platform_Gov_Connectors_Customer_Handout.docx"
PDF_PATH = DOCUMENTS_DIR / "Power_Platform_Gov_Connectors_Customer_Handout.pdf"

OVERVIEW_URL = f"{BASE}/en-us/connectors/connector-reference/"
SOURCE_PAGES = {
    "Power Apps": f"{BASE}/en-us/connectors/connector-reference/connector-reference-powerapps-connectors",
    "Power Automate": f"{BASE}/en-us/connectors/connector-reference/connector-reference-powerautomate-connectors",
}
CLOUDS = ["GCC", "GCC High", "DoD"]
PRODUCT_TITLES = {"Logic Apps", "Power Automate", "Power Apps"}
PRODUCT_ICON_URLS = {
    "Logic Apps": f"{OVERVIEW_URL}media/logicapps-32.png",
    "Power Automate": f"{OVERVIEW_URL}media/powerautomate-32.png",
    "Power Apps": f"{OVERVIEW_URL}media/powerapps-48.png",
}
PRODUCT_ICON_PATHS = {
    product: PRODUCT_ICON_DIR / f"{product.lower().replace(' ', '-')}.png"
    for product in PRODUCT_ICON_URLS
}


def repo_path(path):
    path = Path(path)
    return str(path if path.is_absolute() else REPO_ROOT / path)


def normalize_href(href):
    href = href.split("?", 1)[0].split("#", 1)[0]
    if href.startswith("/"):
        href = BASE + href
    if not href.endswith("/"):
        href += "/"
    return href


def fetch_soup(session, url):
    response = session.get(url, timeout=45)
    response.raise_for_status()
    return BeautifulSoup(response.text, "html.parser"), response.text


def source_updated_at(soup):
    meta = soup.find("meta", attrs={"name": "updated_at"})
    if meta and meta.get("content"):
        return meta["content"]
    meta = soup.find("meta", attrs={"name": "ms.date"})
    return meta["content"] if meta and meta.get("content") else ""


def parse_overview(session):
    soup, _ = fetch_soup(session, OVERVIEW_URL)
    metadata = {}
    for cell in soup.find_all("td"):
        link = cell.find("a")
        images = cell.find_all("img")
        if not link or not images:
            continue
        href = normalize_href(link["href"])
        name = link.get_text(" ", strip=True)
        icon = images[0].get("src", "")
        if icon.startswith("/"):
            icon = BASE + icon

        icon_titles = [(img.get("title") or img.get("alt") or "").strip() for img in images[1:]]
        products = [title for title in icon_titles if title in PRODUCT_TITLES]
        if "MCP Server" in name and "MCP Server" not in products:
            products.append("MCP Server")

        text = cell.get_text("\n", strip=True)
        publisher_match = re.search(r"By:\s*(.+)", text)
        publisher_raw = publisher_match.group(1).strip().rstrip(".") if publisher_match else ""
        publisher_group = "Microsoft" if publisher_raw.lower() == "microsoft" else "Non-Microsoft"
        publisher = publisher_group if not publisher_raw or publisher_group == "Microsoft" else f"{publisher_group}: {publisher_raw}"

        metadata[href] = {
            "name": name,
            "icon_url": icon,
            "products": products,
            "tier": "Premium" if "Premium" in icon_titles else "Standard",
            "status": "Preview" if "Preview" in icon_titles or "(Preview)" in name else "Production",
            "publisher": publisher,
            "publisher_raw": publisher_raw,
            "learn_url": href,
        }
    return metadata, source_updated_at(soup)


def parse_availability(session, product, url):
    soup, _ = fetch_soup(session, url)
    rows = []
    table = soup.find("table")
    for tr in table.find_all("tr")[1:]:
        cells = tr.find_all("td")
        if len(cells) < 5:
            continue
        link = cells[1].find("a")
        if not link:
            continue
        href = normalize_href(link["href"])
        icon = cells[0].find("img")
        availability = {
            "GCC": bool(cells[2].get_text("", strip=True)),
            "GCC High": bool(cells[3].get_text("", strip=True)),
            "DoD": bool(cells[4].get_text("", strip=True)),
        }
        if not any(availability.values()):
            continue
        rows.append(
            {
                "href": href,
                "name": link.get_text(" ", strip=True).replace(" (Preview)", ""),
                "icon_url": icon["src"] if icon and icon.get("src") else "",
                "availability": availability,
                "availability_product": product,
            }
        )
    return rows, source_updated_at(soup)


def build_dataset():
    session = requests.Session()
    session.headers.update({"User-Agent": "Microsoft Scout connector deliverable generator"})
    overview, overview_date = parse_overview(session)
    combined = {}
    source_dates = {OVERVIEW_URL: overview_date}

    for product, url in SOURCE_PAGES.items():
        rows, page_date = parse_availability(session, product, url)
        source_dates[url] = page_date
        for row in rows:
            record = combined.setdefault(
                row["href"],
                {
                    "learn_url": row["href"],
                    "name": row["name"],
                    "icon_url": row["icon_url"],
                    "availability": {cloud: False for cloud in CLOUDS},
                    "availability_products": [],
                },
            )
            for cloud, available in row["availability"].items():
                record["availability"][cloud] = record["availability"][cloud] or available
            if row["availability_product"] not in record["availability_products"]:
                record["availability_products"].append(row["availability_product"])
            if row["icon_url"]:
                record["icon_url"] = row["icon_url"]

    missing_metadata = []
    for href, record in combined.items():
        metadata = overview.get(href)
        if not metadata:
            missing_metadata.append(href)
            continue
        record.update(metadata)
        for product in metadata["products"]:
            if product in ("Power Apps", "Power Automate") and product not in record["availability_products"]:
                record["availability_products"].append(product)

    if missing_metadata:
        raise RuntimeError(f"Missing overview metadata for {len(missing_metadata)} connector(s): {missing_metadata[:5]}")

    records = sorted(combined.values(), key=lambda item: item["name"].lower())
    for record in records:
        record["available_in"] = ", ".join(cloud for cloud in CLOUDS if record["availability"][cloud])
        record["products_display"] = ", ".join(record["products"])
    return records, source_dates


def download_logos(records):
    LOGO_DIR.mkdir(parents=True, exist_ok=True)
    session = requests.Session()
    session.headers.update({"User-Agent": "Microsoft Scout connector deliverable generator"})
    failures = []
    for record in records:
        digest = hashlib.sha1(record["learn_url"].encode("utf-8")).hexdigest()[:14]
        logo_path = LOGO_DIR / f"{digest}.png"
        if logo_path.exists():
            record["logo_path"] = logo_path.relative_to(REPO_ROOT).as_posix()
            continue
        try:
            response = session.get(record["icon_url"], timeout=45)
            response.raise_for_status()
            with Image.open(BytesIO(response.content)) as image:
                image = image.convert("RGBA")
                alpha_bbox = image.getchannel("A").point(lambda value: 255 if value > 12 else 0).getbbox()
                if alpha_bbox:
                    image = image.crop(alpha_bbox)
                white_background = Image.new("RGBA", image.size, (255, 255, 255, 255))
                difference = ImageChops.difference(image, white_background)
                content_bbox = difference.convert("L").point(lambda value: 255 if value > 18 else 0).getbbox()
                if content_bbox:
                    image = image.crop(content_bbox)
                padding = max(4, int(max(image.size) * 0.08))
                padded = Image.new("RGBA", (image.width + padding * 2, image.height + padding * 2), (255, 255, 255, 0))
                padded.alpha_composite(image, (padding, padding))
                image = padded
                image.thumbnail((116, 116), Image.LANCZOS)
                canvas = Image.new("RGBA", (128, 128), (255, 255, 255, 0))
                x = (128 - image.width) // 2
                y = (128 - image.height) // 2
                canvas.alpha_composite(image, (x, y))
                canvas.save(logo_path)
            record["logo_path"] = logo_path.relative_to(REPO_ROOT).as_posix()
        except Exception as exc:
            failures.append((record["name"], record["icon_url"], str(exc)))
    if failures:
        raise RuntimeError(f"Failed to download {len(failures)} logo(s): {failures[:5]}")


def download_product_icons():
    PRODUCT_ICON_DIR.mkdir(parents=True, exist_ok=True)
    session = requests.Session()
    session.headers.update({"User-Agent": "Microsoft Scout connector deliverable generator"})
    for product, url in PRODUCT_ICON_URLS.items():
        path = PRODUCT_ICON_PATHS[product]
        if path.exists():
            continue
        response = session.get(url, timeout=45)
        response.raise_for_status()
        with Image.open(BytesIO(response.content)) as image:
            image.convert("RGBA").save(path)


def add_text(slide, text, x, y, w, h, size=14, color="FFFFFF", bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(w), PptxInches(h))
    frame = box.text_frame
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = PptxPt(size)
    run.font.bold = bold
    run.font.color.rgb = PptxRGBColor.from_string(color)
    return box


@lru_cache(maxsize=None)
def logo_comparison_pixels(path):
    with Image.open(path) as image:
        image = image.convert("RGBA")
        canvas = Image.new("RGBA", image.size, (255, 255, 255, 255))
        canvas.alpha_composite(image)
        canvas = canvas.convert("RGB").resize((64, 64), Image.LANCZOS)
        return canvas.tobytes()


def visually_identical_logo(path_a, path_b):
    pixels_a = logo_comparison_pixels(path_a)
    pixels_b = logo_comparison_pixels(path_b)
    if pixels_a == pixels_b:
        return True
    squared_error = 0
    for left, right in zip(pixels_a, pixels_b):
        difference = left - right
        squared_error += difference * difference
    rms = math.sqrt(squared_error / len(pixels_a))
    return rms <= 1.2


def unique_visual_records(records):
    unique = []
    for record in records:
        if not any(
            visually_identical_logo(repo_path(record["logo_path"]), repo_path(existing["logo_path"]))
            for existing in unique
        ):
            unique.append(record)
    return unique


def add_section(slide, title, records, x, y, w, h, accent):
    slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PptxInches(x),
        PptxInches(y),
        PptxInches(w),
        PptxInches(h),
    ).fill.solid()
    panel = slide.shapes[-1]
    panel.fill.fore_color.rgb = PptxRGBColor(255, 255, 255)
    panel.line.color.rgb = PptxRGBColor.from_string("D9E2F3")

    add_text(slide, title, x + 0.15, y + 0.08, w - 0.3, 0.18, size=8.8, color=accent, bold=True)

    if not records:
        return

    grid_x = x + 0.18
    grid_y = y + 0.34
    grid_w = w - 0.36
    grid_h = h - 0.48
    visible_records = unique_visual_records(records)
    count = len(visible_records)
    gap = 0.035
    best = None
    for candidate_cols in range(8, 31):
        rows = math.ceil(count / candidate_cols)
        if rows > 6:
            continue
        icon_size = min(
            (grid_w - (candidate_cols - 1) * gap) / candidate_cols,
            (grid_h - (rows - 1) * gap) / rows,
            0.36,
        )
        if icon_size <= 0:
            continue
        last_row_count = count - candidate_cols * (rows - 1)
        last_row_fill = last_row_count / candidate_cols
        used_w_ratio = (candidate_cols * icon_size + (candidate_cols - 1) * gap) / grid_w
        orphan_penalty = 0.18 if rows > 1 and last_row_count < max(4, candidate_cols * 0.25) else 0
        score = icon_size + min(last_row_fill, 0.85) * 0.04 - abs(used_w_ratio - 0.82) * 0.025 - orphan_penalty
        if best is None or score > best[0]:
            best = (score, candidate_cols, rows, icon_size)
    if best is None:
        cols = max(1, min(30, math.ceil(math.sqrt(count * grid_w / grid_h))))
        rows = math.ceil(count / cols)
        icon_size = min((grid_w - (cols - 1) * gap) / cols, (grid_h - (rows - 1) * gap) / rows, 0.31)
    else:
        _, cols, rows, icon_size = best
    used_w = cols * icon_size + (cols - 1) * gap
    full_rows = count // cols
    last_row_count = count % cols
    start_x = grid_x + max(0, (grid_w - used_w) / 2)

    for index, record in enumerate(visible_records):
        row = index // cols
        col = index % cols
        row_cols = last_row_count if last_row_count and row == full_rows else cols
        row_used_w = row_cols * icon_size + (row_cols - 1) * gap
        row_start_x = grid_x + max(0, (grid_w - row_used_w) / 2)
        icon_x = row_start_x + col * (icon_size + gap)
        icon_y = grid_y + row * (icon_size + gap)
        slide.shapes.add_picture(
            repo_path(record["logo_path"]),
            PptxInches(icon_x),
            PptxInches(icon_y),
            PptxInches(icon_size),
            PptxInches(icon_size),
        )


def create_pptx(records, source_dates):
    prs = Presentation()
    prs.slide_width = PptxInches(10)
    prs.slide_height = PptxInches(5.625)
    blank = prs.slide_layouts[6]
    colors = {
        "navy": "1E2761",
        "ice": "CADCFC",
        "teal": "028090",
        "mint": "02C39A",
        "charcoal": "36454F",
    }

    for cloud in CLOUDS:
        slide = prs.slides.add_slide(blank)
        background = slide.background
        background.fill.solid()
        background.fill.fore_color.rgb = PptxRGBColor.from_string("F7F9FC")

        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptxInches(0), PptxInches(0), PptxInches(10), PptxInches(0.95)).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = PptxRGBColor.from_string(colors["navy"])
        slide.shapes[-1].line.fill.background()

        cloud_records = [record for record in records if record["availability"][cloud]]
        standard = [record for record in cloud_records if record["tier"] == "Standard"]
        premium = [record for record in cloud_records if record["tier"] == "Premium"]
        standard_unique = unique_visual_records(standard)
        premium_unique = unique_visual_records(premium)
        add_text(slide, f"Power Platform connectors available in {cloud}", 0.35, 0.18, 6.6, 0.35, size=19, color="FFFFFF", bold=True)
        add_text(slide, f"Total connectors: {len(cloud_records)}", 7.05, 0.16, 2.6, 0.38, size=17, color=colors["ice"], bold=True, align=PP_ALIGN.RIGHT)
        add_text(slide, "Microsoft Learn connector reference | Duplicate product logos shown once", 0.35, 0.58, 5.8, 0.2, size=8, color="CADCFC")

        add_section(slide, "STANDARD", standard, 0.45, 1.08, 9.1, 1.76, colors["teal"])
        add_section(slide, "PREMIUM", premium, 0.45, 3.02, 9.1, 1.76, colors["navy"])

        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptxInches(0), PptxInches(5.1), PptxInches(10), PptxInches(0.53)).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = PptxRGBColor.from_string(colors["charcoal"])
        slide.shapes[-1].line.fill.background()
        add_text(slide, f"Total Standard: {len(standard)}", 0.55, 5.25, 3.4, 0.2, size=11.5, color="FFFFFF", bold=True)
        add_text(slide, f"Total Premium: {len(premium)}", 6.05, 5.25, 3.4, 0.2, size=11.5, color="FFFFFF", bold=True, align=PP_ALIGN.RIGHT)
        add_text(slide, f"Icons shown: {len(standard_unique)} standard / {len(premium_unique)} premium", 3.35, 5.24, 3.3, 0.18, size=8.5, color="FFFFFF", align=PP_ALIGN.CENTER)

    prs.save(PPTX_PATH)


def set_cell_text(cell, text, bold=False, size=7.5, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 3 else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Segoe UI Emoji" if any(ord(char) > 127 for char in text) else "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def clear_cell(cell):
    cell.text = ""
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    insert_before = {"w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"}
    inserted = False
    for index, child in enumerate(tc_pr):
        if child.tag in {qn(tag) for tag in insert_before}:
            tc_pr.insert(index, shading)
            inserted = True
            break
    if not inserted:
        tc_pr.append(shading)


def set_cell_borders(cell, **borders):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_element = tc_pr.first_child_found_in("w:tcBorders")
    if borders_element is not None:
        tc_pr.remove(borders_element)
    borders_element = OxmlElement("w:tcBorders")
    insert_before = {
        "w:shd",
        "w:noWrap",
        "w:tcMar",
        "w:textDirection",
        "w:tcFitText",
        "w:vAlign",
        "w:hideMark",
    }
    inserted = False
    for index, child in enumerate(tc_pr):
        if child.tag in {qn(tag) for tag in insert_before}:
            tc_pr.insert(index, borders_element)
            inserted = True
            break
    if not inserted:
        tc_pr.append(borders_element)
    normalized = {}
    for edge, attributes in borders.items():
        normalized[{"left": "start", "right": "end"}.get(edge, edge)] = attributes
    for edge in ["top", "start", "bottom", "end", "insideH", "insideV"]:
        if edge not in normalized:
            continue
        attributes = normalized[edge]
        tag = f"w:{edge}"
        element = borders_element.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders_element.append(element)
        for key, value in attributes.items():
            element.set(qn(f"w:{key}"), str(value))


def mark_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_width(table, width_inches):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.insert(0, width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(int(width_inches * 1440)))


def set_cell_no_wrap(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:noWrap"))
    if existing is not None:
        tc_pr.remove(existing)
    no_wrap = OxmlElement("w:noWrap")
    insert_before = {"w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"}
    inserted = False
    for index, child in enumerate(tc_pr):
        if child.tag in {qn(tag) for tag in insert_before}:
            tc_pr.insert(index, no_wrap)
            inserted = True
            break
    if not inserted:
        tc_pr.append(no_wrap)


def set_cell_margins(cell, top=35, left=35, bottom=35, right=35):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is not None:
        tc_pr.remove(margins)
    margins = OxmlElement("w:tcMar")
    insert_before = {"w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"}
    inserted = False
    for index, child in enumerate(tc_pr):
        if child.tag in {qn(tag) for tag in insert_before}:
            tc_pr.insert(index, margins)
            inserted = True
            break
    if not inserted:
        tc_pr.append(margins)
    for edge, value in {"top": top, "start": left, "bottom": bottom, "end": right}.items():
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cells_middle(cells):
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_picture_run(paragraph, image_path, width, alt_text):
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=width)
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", alt_text)
    return inline_shape


def add_product_icons(cell, products):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for product in ["Logic Apps", "Power Automate", "Power Apps"]:
        if product in products:
            add_picture_run(paragraph, PRODUCT_ICON_PATHS[product], Inches(0.16), product)
            paragraph.add_run(" ")


def publisher_display(record):
    publisher = record.get("publisher_raw") or record.get("publisher", "")
    return publisher.replace("Non-Microsoft: ", "")


def add_hyperlink(paragraph, text, url, size=7.5):
    part = paragraph.part
    relationship_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(int(size * 2)))
    properties.extend([color, underline, size_element])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([properties, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source_links(paragraph):
    add_hyperlink(paragraph, "Connector reference overview", OVERVIEW_URL, size=7.5)
    paragraph.add_run(" | ")
    add_hyperlink(paragraph, "Power Apps connectors", SOURCE_PAGES["Power Apps"], size=7.5)
    paragraph.add_run(" | ")
    add_hyperlink(paragraph, "Power Automate connectors", SOURCE_PAGES["Power Automate"], size=7.5)


def add_key_table(cell):
    clear_cell(cell)
    key_table = cell.add_table(rows=2, cols=3)
    key_table.style = "Table Grid"
    key_widths = [0.88, 0.48, 2.72]
    set_table_width(key_table, sum(key_widths))

    title_cell = key_table.rows[0].cells[0].merge(key_table.rows[0].cells[2])
    title_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade_cell(title_cell, "1E2761")
    set_cell_margins(title_cell, top=16, bottom=16)
    title_para = set_cell_text(title_cell, "KEY", bold=True, size=7.5, color="FFFFFF")
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for index, header in enumerate(["Category", "Icon", "Meaning"]):
        header_cell = key_table.rows[1].cells[index]
        header_cell.width = Inches(key_widths[index])
        header_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(header_cell, "EAF1FB")
        set_cell_margins(header_cell, top=16, bottom=16)
        set_cell_no_wrap(header_cell)
        paragraph = set_cell_text(header_cell, header, bold=True, size=7.0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    rows = [
        ("Tier", "⚫", "Standard", None),
        ("Tier", "💎", "Premium", None),
        ("Status", "🟢", "Production", None),
        ("Status", "💡", "Preview", None),
        ("Cloud", "✅", "Available in that cloud", None),
        ("Product", "", "Logic Apps", PRODUCT_ICON_PATHS["Logic Apps"]),
        ("Product", "", "Power Automate", PRODUCT_ICON_PATHS["Power Automate"]),
        ("Product", "", "Power Apps", PRODUCT_ICON_PATHS["Power Apps"]),
    ]
    for category, icon, meaning, icon_path in rows:
        row = key_table.add_row()
        row.height = Inches(0.18)
        set_cells_middle(row.cells)
        for index, width in enumerate(key_widths):
            row.cells[index].width = Inches(width)
            set_cell_margins(row.cells[index], top=14, bottom=14)
        set_cell_text(row.cells[0], category, size=7.0)
        if icon_path:
            row.cells[1].text = ""
            icon_paragraph = row.cells[1].paragraphs[0]
            icon_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            icon_paragraph.paragraph_format.space_before = Pt(0)
            icon_paragraph.paragraph_format.space_after = Pt(0)
            icon_paragraph.paragraph_format.line_spacing = 1
            add_picture_run(icon_paragraph, icon_path, Inches(0.13), meaning)
        else:
            set_cell_text(row.cells[1], icon, size=7.5)
        set_cell_text(row.cells[2], meaning, size=7.0)

    for start, end, category in [(2, 3, "Tier"), (4, 5, "Status"), (7, 9, "Product")]:
        merged = key_table.rows[start].cells[0].merge(key_table.rows[end].cells[0])
        merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(merged, top=14, bottom=14)
        set_cell_text(merged, category, size=7.0)
    return key_table


def add_box_title(paragraph, text):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)


def add_change_log_table(cell, latest_note):
    clear_cell(cell)
    table = cell.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    set_table_width(table, 3.45)
    table.rows[0].height = Inches(0.24)
    table.rows[1].height = Inches(0.62)
    for row in table.rows:
        set_cells_middle(row.cells)
        set_cell_margins(row.cells[0], top=22, bottom=22, left=60, right=60)
    add_box_title(table.rows[0].cells[0].paragraphs[0], "Change Log")
    paragraph = set_cell_text(table.rows[1].cells[0], latest_note, size=8.0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def add_counts_table(cell, counts, total, clear=True):
    if clear:
        clear_cell(cell)
    table = cell.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    set_table_width(table, 3.45)
    table.rows[0].height = Inches(0.24)
    for row in table.rows:
        set_cells_middle(row.cells)
        for child_cell in row.cells:
            set_cell_margins(child_cell, top=18, bottom=18, left=60, right=60)
    title = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    add_box_title(title.paragraphs[0], "Connector Count")
    count_rows = [
        ("Total", str(total)),
        ("GCC", str(counts["GCC"])),
        ("GCC-High", str(counts["GCC High"])),
        ("DoD", str(counts["DoD"])),
    ]
    for row_index, (label, value) in enumerate(count_rows, start=1):
        left, right = table.rows[row_index].cells
        set_cell_text(left, label, bold=True, size=8.5)
        value_paragraph = set_cell_text(right, value, bold=True, size=8.5)
        value_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return table


def create_docx(records, source_dates, latest_change_note):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.25)
    section.right_margin = Inches(0.25)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Power Platform connectors available in US Government clouds")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string("1E2761")

    counts = {cloud: sum(1 for record in records if record["availability"][cloud]) for cloud in CLOUDS}

    dashboard = document.add_table(rows=1, cols=2)
    dashboard.autofit = False
    set_table_width(dashboard, 8.0)
    dashboard_widths = [4.34, 3.66]
    for index, width in enumerate(dashboard_widths):
        dashboard.rows[0].cells[index].width = Inches(width)
        set_cell_borders(
            dashboard.rows[0].cells[index],
            top={"val": "nil"},
            bottom={"val": "nil"},
            left={"val": "nil"},
            right={"val": "nil"},
        )
        set_cell_margins(dashboard.rows[0].cells[index], top=0, bottom=0, left=0, right=0)
    add_key_table(dashboard.rows[0].cells[0])
    right_cell = dashboard.rows[0].cells[1]
    clear_cell(right_cell)
    add_change_log_table(right_cell, latest_change_note)
    spacer = right_cell.add_paragraph()
    spacer.paragraph_format.space_before = Pt(3)
    spacer.paragraph_format.space_after = Pt(3)
    add_counts_table(right_cell, counts, len(records), clear=False)

    document.add_paragraph().paragraph_format.space_after = Pt(2)

    headers = ["", "Connector", "Tier", "Status", "Product", "Publisher", "GCC", "GCC-High", "DoD"]
    widths = [0.33, 2.05, 0.45, 0.55, 0.85, 2.10, 0.45, 0.75, 0.47]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, sum(widths))
    mark_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        shade_cell(cell, "1E2761")
        set_cell_margins(cell, top=28, bottom=28, left=28, right=28)
        set_cell_no_wrap(cell)
        paragraph = set_cell_text(cell, header, bold=True, size=7.0, color="FFFFFF")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_borders(table.rows[0].cells[0], right={"val": "nil"})
    set_cell_borders(table.rows[0].cells[1], left={"val": "nil"})

    for record in records:
        row = table.add_row()
        row.height = Inches(0.24)
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        logo_cell = row.cells[0]
        logo_cell.text = ""
        logo_paragraph = logo_cell.paragraphs[0]
        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_picture_run(logo_paragraph, repo_path(record["logo_path"]), Inches(0.17), record["name"])
        set_cell_borders(logo_cell, right={"val": "nil"})

        connector_cell = row.cells[1]
        connector_cell.text = ""
        add_hyperlink(connector_cell.paragraphs[0], record["name"], record["learn_url"], size=7.0)
        set_cell_borders(connector_cell, left={"val": "nil"})
        set_cell_text(row.cells[2], "💎" if record["tier"] == "Premium" else "⚫", size=7.5)
        set_cell_text(row.cells[3], "💡" if record["status"] == "Preview" else "🟢", size=7.5)
        add_product_icons(row.cells[4], record["products"])
        set_cell_text(row.cells[5], publisher_display(record), size=6.8)
        for cloud, cell_index in [("GCC", 6), ("GCC High", 7), ("DoD", 8)]:
            set_cell_text(row.cells[cell_index], "✅" if record["availability"][cloud] else "", size=7.5)

    sources = document.add_paragraph()
    sources.paragraph_format.space_before = Pt(4)
    sources.paragraph_format.space_after = Pt(0)
    sources.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = sources.add_run("Sources: ")
    label.bold = True
    label.font.name = "Arial"
    label.font.size = Pt(7.5)
    add_source_links(sources)

    document.save(DOCX_PATH)
    repair_docx_schema(DOCX_PATH)


def repair_docx_schema(path):
    temp_dir = Path(tempfile.mkdtemp(prefix="docx-repair-"))
    try:
        with zipfile.ZipFile(path, "r") as archive:
            archive.extractall(temp_dir)

        document_xml = temp_dir / "word" / "document.xml"
        text = document_xml.read_text(encoding="utf-8")
        text = re.sub(r"<w:shd(?![^>]w:val=)", '<w:shd w:val="clear"', text)
        document_xml.write_text(text, encoding="utf-8")

        settings_xml = temp_dir / "word" / "settings.xml"
        if settings_xml.exists():
            settings = settings_xml.read_text(encoding="utf-8")
            settings = re.sub(r"<w:zoom(?![^>]w:percent=)([^>]*)/>", r'<w:zoom\1 w:percent="100"/>', settings)
            settings = re.sub(
                r'(<w:compatSetting w:name="compatibilityMode"[^>]*w:val=")14("[^>]*/>)',
                r"\g<1>15\2",
                settings,
            )
            settings_xml.write_text(settings, encoding="utf-8")

        font_table_xml = temp_dir / "word" / "fontTable.xml"
        if font_table_xml.exists():
            font_table = font_table_xml.read_text(encoding="utf-8")
            font_table = re.sub(r"\s*<w:font w:name=\"[^\"]*[^\x00-\x7F][^\"]*\">.*?</w:font>", "", font_table, flags=re.DOTALL)
            font_table_xml.write_text(font_table, encoding="utf-8")

        repaired = path.with_suffix(".repaired.docx")
        if repaired.exists():
            repaired.unlink()
        with zipfile.ZipFile(repaired, "w", zipfile.ZIP_DEFLATED) as archive:
            for file_path in temp_dir.rglob("*"):
                if file_path.is_file():
                    archive.write(file_path, file_path.relative_to(temp_dir).as_posix())
        repaired.replace(path)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def convert_docx_to_pdf():
    program_files = Path(os.environ.get("PROGRAMFILES", ""))
    candidates = [
        shutil.which("libreoffice"),
        shutil.which("soffice"),
        program_files / "LibreOffice" / "program" / "soffice.exe",
    ]
    executable = next((str(candidate) for candidate in candidates if candidate and Path(candidate).is_file()), None)
    if executable is None:
        raise RuntimeError(
            "LibreOffice is required to generate the PDF. Install LibreOffice and ensure "
            "'libreoffice' or 'soffice' is available on PATH."
        )

    PDF_PATH.unlink(missing_ok=True)
    result = subprocess.run(
        [
            executable,
            "--headless",
            "--convert-to",
            "pdf:writer_pdf_Export",
            "--outdir",
            str(DOCUMENTS_DIR),
            str(DOCX_PATH),
        ],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    if result.returncode != 0 or not PDF_PATH.is_file() or PDF_PATH.stat().st_size == 0:
        output = "\n".join(part.strip() for part in (result.stdout, result.stderr) if part.strip())
        raise RuntimeError(f"LibreOffice failed to generate {PDF_PATH.name}.\n{output}")


def validate_outputs():
    required_paths = [DATA_PATH, CHANGE_LOG_PATH, PPTX_PATH, DOCX_PATH, PDF_PATH]
    invalid_paths = [path for path in required_paths if not path.is_file() or path.stat().st_size == 0]
    if invalid_paths:
        raise RuntimeError(f"Missing or empty generated output(s): {', '.join(map(str, invalid_paths))}")

    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    if not isinstance(data.get("records"), list) or not data["records"]:
        raise RuntimeError(f"{DATA_PATH} does not contain a non-empty records list.")


def load_previous_records():
    if not DATA_PATH.exists():
        return []
    try:
        previous = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []
    return previous.get("records", [])


def load_change_log():
    if not CHANGE_LOG_PATH.exists():
        return []
    try:
        return json.loads(CHANGE_LOG_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []


def diff_connectors(previous_records, new_records):
    """Compare two dataset snapshots and describe what changed, mirroring the
    manually-authored entries that used to live in the change log."""
    previous_by_url = {record["learn_url"]: record for record in previous_records}
    new_by_url = {record["learn_url"]: record for record in new_records}

    added = removed = cloud_changes = metadata_changes = 0
    details = []

    for url, record in new_by_url.items():
        if url in previous_by_url:
            continue
        added += 1
        details.append(
            f"Added connector: {record['name']} ({record['tier']}, {record['status']}; {record['available_in']})."
        )

    for url, record in previous_by_url.items():
        if url in new_by_url:
            continue
        removed += 1
        details.append(
            f"Removed connector: {record['name']} (previously {record['tier']}, {record['status']}; {record['available_in']})."
        )

    for url, new_record in new_by_url.items():
        old_record = previous_by_url.get(url)
        if old_record is None:
            continue
        if old_record["name"] != new_record["name"]:
            metadata_changes += 1
            details.append(f"Renamed connector: {old_record['name']} \u2192 {new_record['name']}.")
        if old_record["status"] != new_record["status"]:
            metadata_changes += 1
            details.append(
                f"Changed release status for {new_record['name']}: {old_record['status']} \u2192 {new_record['status']}."
            )
        if old_record["tier"] != new_record["tier"]:
            metadata_changes += 1
            details.append(
                f"Changed tier for {new_record['name']}: {old_record['tier']} \u2192 {new_record['tier']}."
            )
        for cloud in CLOUDS:
            old_available = old_record["availability"].get(cloud)
            new_available = new_record["availability"].get(cloud)
            if old_available == new_available:
                continue
            cloud_changes += 1
            verb = "Added" if new_available else "Removed"
            details.append(f"{verb} {cloud} availability for {new_record['name']}.")

    counts = {
        "added": added,
        "removed": removed,
        "cloud_changes": cloud_changes,
        "metadata_changes": metadata_changes,
    }
    return counts, details


def summarize_changes(counts, details):
    def plural(count, noun):
        return f"{count} {noun}{'' if count == 1 else 's'}"

    headline_parts = []
    if counts["added"]:
        headline_parts.append(f"{plural(counts['added'], 'connector')} added")
    if counts["removed"]:
        headline_parts.append(f"{plural(counts['removed'], 'connector')} removed")
    if counts["cloud_changes"]:
        headline_parts.append(f"{plural(counts['cloud_changes'], 'cloud availability change')}")
    if counts["metadata_changes"]:
        headline_parts.append(f"{plural(counts['metadata_changes'], 'metadata/product change')}")

    if not headline_parts:
        return "No changes detected. Regenerated files from the current Microsoft Learn data."

    headline = "; ".join(headline_parts) + "."
    return " ".join([headline, *details])


def record_change_log_entry(previous_records, new_records):
    counts, details = diff_connectors(previous_records, new_records)
    if not any(counts.values()):
        return None

    entry = {
        "date": datetime.now(timezone.utc).strftime("%B %d, %Y"),
        "summary": summarize_changes(counts, details),
        "details": details,
    }
    change_log = load_change_log()
    change_log.append(entry)
    CHANGE_LOG_PATH.write_text(json.dumps(change_log, indent=2), encoding="utf-8")
    return entry


def latest_change_note(change_log):
    if not change_log:
        return "No changes recorded yet."
    entry = change_log[-1]
    note = f"{entry['date']}: {entry['summary']}"
    if len(note) > 220:
        note = note[:217].rstrip() + "..."
    return note


def main():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)

    previous_records = load_previous_records()
    records, source_dates = build_dataset()
    download_logos(records)
    download_product_icons()

    new_entry = record_change_log_entry(previous_records, records)

    dataset_json = json.dumps({"source_dates": source_dates, "records": records}, indent=2)
    previous_dataset_json = DATA_PATH.read_text(encoding="utf-8") if DATA_PATH.exists() else None
    dataset_changed = dataset_json != previous_dataset_json
    DATA_PATH.write_text(dataset_json, encoding="utf-8")

    document_paths = [PPTX_PATH, DOCX_PATH, PDF_PATH]
    documents_regenerated = dataset_changed or any(
        not path.is_file() or path.stat().st_size == 0 for path in document_paths
    )
    if documents_regenerated:
        create_pptx(records, source_dates)
        create_docx(records, source_dates, latest_change_note(load_change_log()))
        convert_docx_to_pdf()
    validate_outputs()

    counts = {
        cloud: {
            "total": sum(1 for record in records if record["availability"][cloud]),
            "standard": sum(1 for record in records if record["availability"][cloud] and record["tier"] == "Standard"),
            "premium": sum(1 for record in records if record["availability"][cloud] and record["tier"] == "Premium"),
            "standard_icons_shown": len(unique_visual_records([record for record in records if record["availability"][cloud] and record["tier"] == "Standard"])),
            "premium_icons_shown": len(unique_visual_records([record for record in records if record["availability"][cloud] and record["tier"] == "Premium"])),
        }
        for cloud in CLOUDS
    }
    print(
        json.dumps(
            {
                "records": len(records),
                "counts": counts,
                "pptx": str(PPTX_PATH),
                "docx": str(DOCX_PATH),
                "pdf": str(PDF_PATH),
                "documents_regenerated": documents_regenerated,
                "change_log_entry": new_entry,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
